"""Conversion engines: images, audio, video, PDF/TXT/DOCX, documents, archives."""

from __future__ import annotations

import gzip
import io
import os
import shutil
import subprocess
import tarfile
import tempfile
import threading
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from PIL import Image, ImageOps

from . import naming
from .catalog import (
    AUDIO_EXTS, FAMILY_ARCHIVE, FAMILY_AUDIO, FAMILY_DOC, FAMILY_GIF,
    FAMILY_IMAGE, FAMILY_PDF, FAMILY_TXT, FAMILY_VIDEO, GIF_EXTS, PDF_EXTS,
    RASTER_EXTS, SVG_EXTS, TXT_EXTS, ARCHIVE_EXTS, family_of,
)
from .media import (
    CREATE_NO_WINDOW, MediaInfo, ffmpeg_path, probe, run_ffmpeg, run_process,
    unrar_path, rar_path,
)

try:
    import pillow_heif

    pillow_heif.register_heif_opener()
except Exception:
    pillow_heif = None


class EngineError(Exception):
    pass


_MAX_EXTRACT_BYTES = 16 * 1024 ** 3
_GZ_SNIFF_BYTES = 64 * 1024


@dataclass
class Ctx:
    """Progress/cancel/status plumbing shared by every conversion."""

    progress: Callable[[float], None] = lambda f: None
    status: Callable[[str], None] = lambda s: None
    cancel: threading.Event | None = None

    def cancelled(self) -> bool:
        return self.cancel is not None and self.cancel.is_set()


def _no_window_kw() -> dict:
    return {"creationflags": CREATE_NO_WINDOW} if os.name == "nt" else {}


def verify_writable(path: Path) -> None:
    """Reject locked files early (macOS parity: locked-file rejection)."""
    if not path.exists():
        raise EngineError(f"The file no longer exists: {path.name}")
    try:
        with open(path, "rb+"):
            pass
    except PermissionError as exc:
        raise EngineError(f"The file is locked or read-only: {path.name}") from exc
    except OSError:
        return


def _unique(folder: Path, stem: str, suffix: str, reserved: set[Path]) -> Path:
    return naming.unique_path(folder, stem, suffix, reserved)


# ---------------------------------------------------------------------------
# Images
# ---------------------------------------------------------------------------

def load_image(path: Path) -> Image.Image:
    if path.suffix.lower() in SVG_EXTS:
        return render_svg(path)
    try:
        img = Image.open(path)
        img.load()
        return img
    except Exception as exc:
        raise EngineError(f"Could not read image: {path.name}\n{exc}") from exc


def render_svg(path: Path, target_width: int | None = None) -> Image.Image:
    """Rasterize SVG at its declared pixel size via QtSvg."""
    from PySide6.QtCore import QByteArray, QSize
    from PySide6.QtGui import QImage, QPainter
    from PySide6.QtSvg import QSvgRenderer

    data = path.read_bytes()
    if path.suffix.lower() == ".svgz":
        data = gzip.decompress(data)
    renderer = QSvgRenderer(QByteArray(data))
    if not renderer.isValid():
        raise EngineError(f"Could not rasterize SVG: {path.name}")
    size = renderer.defaultSize()
    width = target_width or (size.width() if size.width() > 0 else 1024)
    height = int(round(width * size.height() / size.width())) if size.width() else width
    if height <= 0:
        height = 1024
    image = QImage(width, height, QImage.Format_ARGB32_Premultiplied)
    image.fill(0)
    painter = QPainter(image)
    renderer.render(painter)
    painter.end()
    buffer = image.constBits().tobytes()
    return Image.frombuffer("RGBA", (width, height), buffer, "raw", "BGRA", image.bytesPerLine(), 1)


def _flatten(img: Image.Image, background: str = "white") -> Image.Image:
    if img.mode in ("RGBA", "LA", "P"):
        img = img.convert("RGBA")
        canvas = Image.new("RGB", img.size, background)
        canvas.paste(img, mask=img.getchannel("A"))
        return canvas
    return img.convert("RGB")


def _image_save_args(target: str, img: Image.Image, source: Path) -> tuple[dict, str]:
    kwargs: dict = {}
    fmt = target.upper()
    if target == "jpg":
        fmt = "JPEG"
        kwargs["quality"] = 92
        kwargs["optimize"] = True
        kwargs["progressive"] = True
    elif target == "png":
        kwargs["optimize"] = True
    elif target == "webp":
        fmt = "WEBP"
        kwargs["quality"] = 92
        kwargs["method"] = 6
    elif target == "heic":
        fmt = "HEIF"
        kwargs["quality"] = 90
    elif target == "tiff":
        fmt = "TIFF"
        kwargs["compression"] = "tiff_lzw"
    return kwargs, fmt


def convert_image(path: Path, target: str, ctx: Ctx, reserved: set[Path]) -> Path:
    verify_writable(path)
    ctx.status(f"Converting {path.name} to {target.upper()}")
    img = load_image(path)
    folder = path.parent
    out_path = _unique(folder, path.stem, f".{target}", reserved)

    img = ImageOps.exif_transpose(img)
    exif = None
    if target.lower() not in ("png", "heic", "heif"):
        exif = img.info.get("exif")

    if target == "pdf":
        if img.mode in ("RGBA", "LA", "P"):
            img = _flatten(img)
        elif img.mode != "RGB":
            img = img.convert("RGB")
        img.save(out_path, "PDF", resolution=150.0)
        ctx.progress(1.0)
        return out_path

    if target in ("jpg",):
        img = _flatten(img)
    kwargs, fmt = _image_save_args(target, img, path)
    if exif and fmt in ("JPEG", "TIFF", "WEBP"):
        kwargs["exif"] = exif
    try:
        if fmt == "JPEG" and img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        img.save(out_path, fmt, **kwargs)
    except Exception as exc:
        raise EngineError(f"Could not save {out_path.name}\n{exc}") from exc
    ctx.progress(1.0)
    return out_path


def images_to_pdf(paths: list[Path], ctx: Ctx, reserved: set[Path]) -> Path:
    first = paths[0]
    folder = first.parent
    name = naming.folder_name(first, "PDF") if len(paths) > 1 else first.stem
    out_path = _unique(folder, name, ".pdf", reserved)
    pages: list[Image.Image] = []
    try:
        for path in paths:
            verify_writable(path)
            try:
                img = Image.open(path)
            except Exception as exc:
                raise EngineError(f"Could not read image: {path.name}\n{exc}") from exc
            if img.getexif().get(274, 1) != 1:
                img = ImageOps.exif_transpose(img)
            if img.mode in ("RGBA", "LA", "P"):
                img = _flatten(img)
            elif img.mode != "RGB":
                img = img.convert("RGB")
            pages.append(img)
        total = len(pages)
        if total == 0:
            raise EngineError("No images selected.")
        pages[0].save(
            out_path, "PDF", resolution=150.0,
            save_all=True, append_images=pages[1:],
        )
    finally:
        for page in pages:
            try:
                page.close()
            except Exception:
                pass
    ctx.progress(1.0)
    return out_path


def image_to_docx(path: Path, ctx: Ctx, reserved: set[Path]) -> Path:
    import docx
    from docx.shared import Inches

    verify_writable(path)
    out_path = _unique(path.parent, path.stem, ".docx", reserved)
    document = docx.Document()
    img = ImageOps.exif_transpose(load_image(path))
    tmp = io.BytesIO()
    (img.convert("RGB") if img.mode != "RGB" else img).save(tmp, "PNG")
    tmp.seek(0)
    width_inches = min(6.5, img.width / 150.0)
    document.add_picture(tmp, width=Inches(width_inches))
    document.save(str(out_path))
    ctx.progress(1.0)
    return out_path


# ---------------------------------------------------------------------------
# Audio / video via FFmpeg
# ---------------------------------------------------------------------------

CODEC_ARGS = {
    "mp3": ["-c:a", "libmp3lame", "-q:a", "2"],
    "m4a": ["-c:a", "aac", "-b:a", "192k"],
    "wav": ["-c:a", "pcm_s16le"],
    "flac": ["-c:a", "flac", "-compression_level", "6"],
}


def convert_audio(path: Path, target: str, ctx: Ctx, reserved: set[Path]) -> Path:
    verify_writable(path)
    info = probe(path)
    out_path = _unique(path.parent, path.stem, f".{target}", reserved)
    ctx.status(f"Converting {path.name} to {target.upper()}")
    args = ["-i", str(path), "-vn", *CODEC_ARGS[target], "-map_metadata", "0", str(out_path)]
    tail: list[str] = []
    code = run_ffmpeg(args, info.duration, ctx.progress, ctx.cancel, tail=tail)
    _finish_ffmpeg(code, out_path, tail)
    return out_path


_EVEN_SCALE = "scale=trunc(iw/2)*2:trunc(ih/2)*2"


def convert_video(path: Path, target: str, ctx: Ctx, reserved: set[Path]) -> Path:
    verify_writable(path)
    info = probe(path)
    out_path = _unique(path.parent, path.stem, f".{target}", reserved)
    ctx.status(f"Converting {path.name} to {target.upper()}")
    args = ["-i", str(path)]
    if target == "mp4":
        args += [
            "-vf", _EVEN_SCALE,
            "-c:v", "libx264", "-crf", "20", "-preset", "veryfast",
            "-pix_fmt", "yuv420p", "-movflags", "+faststart",
            "-c:a", "aac", "-b:a", "192k",
        ]
    elif target == "mov":
        args += [
            "-vf", _EVEN_SCALE,
            "-c:v", "libx264", "-crf", "20", "-preset", "veryfast",
            "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "192k",
        ]
    elif target == "mkv":
        codecs_copy = info.video() and info.video().codec in ("h264", "hevc", "vp9", "av1")
        if codecs_copy:
            args += ["-c:v", "copy"]
        else:
            args += ["-vf", _EVEN_SCALE, "-c:v", "libx264", "-crf", "20", "-preset", "veryfast", "-pix_fmt", "yuv420p"]
        audio = info.audio()
        if audio and audio.codec in ("aac", "mp3", "flac", "opus", "vorbis"):
            args += ["-c:a", "copy"]
        else:
            args += ["-c:a", "aac", "-b:a", "192k"]
    elif target == "gif":
        args += ["-vf", _gif_filter(), "-loop", "0"]
    elif target in ("mp3", "m4a"):
        args += ["-vn", *CODEC_ARGS[target]]
    args += [str(out_path)]
    tail: list[str] = []
    code = run_ffmpeg(args, info.duration, ctx.progress, ctx.cancel, tail=tail)
    _finish_ffmpeg(code, out_path, tail)
    return out_path


def _gif_filter() -> str:
    return (
        "fps=15,scale='min(960,iw)':-2:flags=lanczos,"
        "split[a][b];[a]palettegen=stats_mode=diff[p];[b][p]paletteuse=dither=bayer:bayer_scale=3"
    )


def convert_gif(path: Path, target: str, ctx: Ctx, reserved: set[Path]) -> Path:
    verify_writable(path)
    info = probe(path)
    out_path = _unique(path.parent, path.stem, f".{target}", reserved)
    ctx.status(f"Converting {path.name} to {target.upper()}")
    args = ["-i", str(path)]
    if target == "mp4":
        args += ["-vf", _EVEN_SCALE, "-c:v", "libx264", "-crf", "20", "-preset", "veryfast",
                 "-pix_fmt", "yuv420p", "-movflags", "+faststart", "-an"]
    elif target == "mov":
        args += ["-vf", _EVEN_SCALE, "-c:v", "libx264", "-crf", "20", "-preset", "veryfast",
                 "-pix_fmt", "yuv420p", "-an"]
    else:
        args += ["-vf", _EVEN_SCALE, "-c:v", "libx264", "-crf", "20", "-preset", "veryfast",
                 "-pix_fmt", "yuv420p"]
    args += [str(out_path)]
    tail: list[str] = []
    code = run_ffmpeg(args, info.duration, ctx.progress, ctx.cancel, tail=tail)
    _finish_ffmpeg(code, out_path, tail)
    return out_path


def _finish_ffmpeg(code: int, out_path: Path, tail: list[str] | None = None) -> None:
    if code == -9:
        if out_path.exists():
            try:
                out_path.unlink()
            except OSError:
                pass
        raise EngineError("Cancelled.")
    if code != 0 or not out_path.exists():
        out_path.unlink(missing_ok=True)
        detail = ""
        if tail:
            lines = [line for line in tail if line][-3:]
            if lines:
                detail = ": " + " | ".join(lines)
        raise EngineError(f"FFmpeg could not process {out_path.name}{detail}")
    if out_path.stat().st_size == 0:
        out_path.unlink(missing_ok=True)
        raise EngineError("FFmpeg produced an empty file.")


# ---------------------------------------------------------------------------
# PDF / TXT / DOCX
# ---------------------------------------------------------------------------

_PDF_DPI = 300.0


def _open_pdf_reader(path: Path):
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise EngineError("This PDF could not be opened.") from exc
    if reader.is_encrypted:
        raise EngineError("Password-protected PDFs are not supported.")
    return reader


def _open_pdf_document(path: Path):
    import pypdfium2 as pdfium

    try:
        return pdfium.PdfDocument(str(path))
    except Exception as exc:
        if "password" in str(exc).lower():
            raise EngineError("Password-protected PDFs are not supported.") from exc
        raise EngineError("This PDF could not be opened.") from exc


def pdf_to_images(path: Path, target: str, ctx: Ctx, reserved: set[Path]) -> list[Path]:
    verify_writable(path)
    pdf = _open_pdf_document(path)
    return _render_pdf_pages(pdf, path.parent, path, target, ctx, reserved)


def _render_pdf_pages(
    pdf, folder: Path, naming_source: Path, target: str, ctx: Ctx, reserved: set[Path],
) -> list[Path]:
    """Rasterize an open pdfium document, named after *naming_source*.

    Single page: ``<stem>.<ext>`` inside *folder*. Multiple pages: a staged
    ``<stem> JPG/PNG Pages`` folder published atomically into *folder*.
    """
    page_count = len(pdf)
    scale = _PDF_DPI / 72.0
    ext = ".jpg" if target == "jpg" else ".png"
    outputs: list[Path] = []
    if page_count <= 1:
        page = pdf[0]
        image = page.render(scale=scale).to_pil()
        out_path = _unique(folder, naming_source.stem, ext, reserved)
        _save_rendered(image, out_path, target)
        outputs.append(out_path)
        ctx.progress(1.0)
        return outputs
    dir_name = naming.folder_name(naming_source, "JPG Pages" if target == "jpg" else "PNG Pages")
    out_dir = folder / dir_name
    index = 2
    while out_dir.exists():
        out_dir = folder / f"{dir_name} {index}"
        index += 1
    tmp_dir = Path(tempfile.mkdtemp(prefix=".tangerine_pages_", dir=str(folder)))
    try:
        for number in range(page_count):
            if ctx.cancelled():
                raise EngineError("Cancelled.")
            page = pdf[number]
            image = page.render(scale=scale).to_pil()
            page_name = f"Page {number + 1:03d}{ext}"
            _save_rendered(image, tmp_dir / page_name, target)
            ctx.progress((number + 1) / page_count)
            ctx.status(f"Rendering page {number + 1} of {page_count}")
        os.replace(tmp_dir, out_dir)
    except Exception:
        shutil.rmtree(tmp_dir, ignore_errors=True)
        raise
    outputs.extend(sorted(out_dir.iterdir()))
    return outputs


def _save_rendered(image, out_path: Path, target: str) -> None:
    if target == "jpg":
        if image.mode != "RGB":
            image = image.convert("RGB")
        image.save(out_path, "JPEG", quality=92, dpi=(_PDF_DPI, _PDF_DPI), optimize=True)
    else:
        image.save(out_path, "PNG", dpi=(_PDF_DPI, _PDF_DPI), optimize=True)


def pdf_text(path: Path, ctx: Ctx | None = None) -> str:
    reader = _open_pdf_reader(path)
    chunks = []
    for number, page in enumerate(reader.pages, start=1):
        if ctx is not None and ctx.cancelled():
            raise EngineError("Cancelled.")
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if len(reader.pages) > 1:
            chunks.append(f"----- Page {number} -----\n{text}")
        else:
            chunks.append(text)
    result = "\n\n".join(chunks).strip()
    if result:
        return result
    try:
        pdf = _open_pdf_document(path)
        pages = []
        for number in range(len(pdf)):
            if ctx is not None and ctx.cancelled():
                raise EngineError("Cancelled.")
            page = pdf[number]
            text = page.get_textpage().get_text_range() or ""
            if len(pdf) > 1:
                pages.append(f"----- Page {number + 1} -----\n{text}")
            else:
                pages.append(text)
        return "\n\n".join(pages).strip()
    except EngineError:
        raise
    except Exception:
        return ""


def convert_pdf(path: Path, target: str, ctx: Ctx, reserved: set[Path]) -> list[Path]:
    verify_writable(path)
    if target in ("jpg", "png"):
        return pdf_to_images(path, target, ctx, reserved)
    if target == "txt":
        ctx.status(f"Extracting text from {path.name}")
        text = pdf_text(path, ctx)
        if not text:
            raise EngineError(
                "This PDF has no selectable text, so OCR would be required.\n"
                "Tangerine's text export does not perform OCR."
            )
        out_path = _unique(path.parent, path.stem, ".txt", reserved)
        out_path.write_text(text + "\n", "utf-8")
        ctx.progress(1.0)
        return [out_path]
    if target == "docx":
        return [pdf_to_docx(path, ctx, reserved)]
    raise EngineError(f"Unsupported PDF conversion: {target}")


def pdf_to_docx(path: Path, ctx: Ctx, reserved: set[Path]) -> Path:
    import docx
    from docx.shared import Inches

    reader = _open_pdf_reader(path)
    document = docx.Document()
    pdf = _open_pdf_document(path)
    scale = 150.0 / 72.0
    for number, page in enumerate(reader.pages):
        if ctx.cancelled():
            raise EngineError("Cancelled.")
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            for paragraph in text.splitlines():
                document.add_paragraph(paragraph)
        else:
            rendered = pdf[number].render(scale=scale).to_pil()
            if rendered.mode != "RGB":
                rendered = rendered.convert("RGB")
            tmp = io.BytesIO()
            rendered.save(tmp, "PNG")
            tmp.seek(0)
            document.add_picture(tmp, width=Inches(6.5))
        ctx.progress((number + 1) / max(len(reader.pages), 1))
    out_path = _unique(path.parent, path.stem, ".docx", reserved)
    document.save(str(out_path))
    return out_path


FONT_CANDIDATES = [
    Path("C:/Windows/Fonts/consola.ttf"),
    Path("C:/Windows/Fonts/cour.ttf"),
    Path("C:/Windows/Fonts/lucon.ttf"),
]

PDF_MONO_CANDIDATES = [
    Path("C:/Windows/Fonts/consola.ttf"),
    Path("C:/Windows/Fonts/DejaVuSansMono.ttf"),
]

PDF_UNICODE_FALLBACKS = [
    Path("C:/Windows/Fonts/msyh.ttc"),
    Path("C:/Windows/Fonts/simsun.ttc"),
]


def _mono_font(size: int):
    from PIL import ImageFont

    for candidate in FONT_CANDIDATES:
        if candidate.exists():
            try:
                return ImageFont.truetype(str(candidate), size)
            except OSError:
                continue
    return ImageFont.load_default()


def _pdf_mono_font(text: str) -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    required = {char for char in text if ord(char) > 0xFF}
    chosen = None
    for index, candidate in enumerate(PDF_MONO_CANDIDATES + PDF_UNICODE_FALLBACKS):
        if not candidate.exists():
            continue
        name = "TangerineMono" if index == 0 else f"TangerineMono{index + 1}"
        try:
            font = TTFont(name, str(candidate), subfontIndex=0)
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(font)
        except Exception:
            continue
        if chosen is None:
            chosen = name
        if all(ord(char) in font.face.charToGlyph for char in required):
            return name
    return chosen or "Courier"


def text_to_pdf(text: str, out_path: Path, ctx: Ctx) -> Path:
    """Write plain text to *out_path*: monospace, 54pt margins, wrapped to fit."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfgen import canvas as pdf_canvas

    ctx.status(f"Writing {out_path.name}")
    margin = 54.0
    font_size = 10.0
    leading = 13.5
    font_name = "Courier"
    try:
        text.encode("latin-1")
    except UnicodeEncodeError:
        font_name = _pdf_mono_font(text)
    width, height = letter
    char_width = pdfmetrics.stringWidth("M", font_name, font_size)
    max_chars = max(20, int((width - 2 * margin) / char_width)) if char_width else 84
    wrapped: list[str] = []
    for raw_line in text.replace("\t", "    ").splitlines():
        line = raw_line
        while len(line) > max_chars:
            wrapped.append(line[:max_chars])
            line = line[max_chars:]
        wrapped.append(line)
    pdf = pdf_canvas.Canvas(str(out_path), pagesize=letter)
    pdf.setFont(font_name, font_size)
    y = height - margin
    for line in wrapped:
        if ctx.cancelled():
            pdf.save()
            out_path.unlink(missing_ok=True)
            raise EngineError("Cancelled.")
        if y < margin:
            pdf.showPage()
            pdf.setFont(font_name, font_size)
            y = height - margin
        pdf.drawString(margin, y, line)
        y -= leading
    pdf.save()
    ctx.progress(1.0)
    return out_path


def txt_to_pdf(path: Path, ctx: Ctx, reserved: set[Path]) -> Path:
    text = path.read_text("utf-8", errors="replace")
    out_path = _unique(path.parent, path.stem, ".pdf", reserved)
    return text_to_pdf(text, out_path, ctx)


def txt_to_image(path: Path, target: str, ctx: Ctx, reserved: set[Path]) -> list[Path]:
    from PIL import ImageDraw

    text = path.read_text("utf-8", errors="replace")
    font = _mono_font(15)
    line_height = 20
    max_chars = 100
    lines: list[str] = []
    for raw_line in text.splitlines():
        line = raw_line.replace("\t", "    ")
        while len(line) > max_chars:
            lines.append(line[:max_chars])
            line = line[max_chars:]
        lines.append(line)
    width = 1275
    ext = ".jpg" if target == "jpg" else ".png"
    if len(lines) * line_height > 16000:
        chunks = [lines[index:index + 800] for index in range(0, len(lines), 800)]
    else:
        chunks = [lines]
    outputs: list[Path] = []
    for chunk in chunks:
        height = max(len(chunk) * line_height + 40, 200)
        image = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(image)
        y = 20
        for line in chunk:
            draw.text((20, y), line, font=font, fill="black")
            y += line_height
        out_path = _unique(path.parent, path.stem, ext, reserved)
        if target == "jpg":
            image.save(out_path, "JPEG", quality=92, optimize=True)
        else:
            image.save(out_path, "PNG", optimize=True)
        outputs.append(out_path)
    ctx.progress(1.0)
    return outputs


def convert_txt(path: Path, target: str, ctx: Ctx, reserved: set[Path]) -> list[Path]:
    verify_writable(path)
    ctx.status(f"Converting {path.name} to {target.upper()}")
    if target == "pdf":
        return [txt_to_pdf(path, ctx, reserved)]
    return txt_to_image(path, target, ctx, reserved)


# ---------------------------------------------------------------------------
# Archives
# ---------------------------------------------------------------------------

def _check_extract_size(total: int) -> None:
    if total > _MAX_EXTRACT_BYTES:
        raise EngineError("This archive is too large to extract safely.")


def _extract_tar(archive: tarfile.TarFile, dest: Path) -> None:
    try:
        archive.extractall(dest, filter="data")
    except TypeError:
        archive.extractall(dest)


def _looks_like_tar(data: bytes) -> bool:
    return len(data) >= 262 and data[257:262] == b"ustar"


def _extract_archive(path: Path, dest: Path, ctx: Ctx) -> None:
    suffix = path.suffix.lower()
    if suffix == ".zip":
        with zipfile.ZipFile(path) as archive:
            _check_extract_size(sum(info.file_size for info in archive.infolist()))
            if ctx.cancelled():
                raise EngineError("Cancelled.")
            archive.extractall(dest)
        return
    if suffix == ".rar":
        unrar = unrar_path()
        if not unrar:
            raise EngineError("WinRAR is required to read RAR archives.")
        if ctx.cancelled():
            raise EngineError("Cancelled.")
        code = run_process([unrar, "x", "-y", str(path), str(dest) + os.sep], cancel=ctx.cancel)
        if code == -9:
            raise EngineError("Cancelled.")
        if code != 0:
            raise EngineError("Could not extract this RAR archive (password-protected or multivolume RAR is not supported).")
        return
    if suffix == ".tar":
        with tarfile.open(path, "r:") as archive:
            _check_extract_size(sum(member.size for member in archive.getmembers()))
            if ctx.cancelled():
                raise EngineError("Cancelled.")
            _extract_tar(archive, dest)
        return
    if suffix == ".gz":
        _extract_gzip(path, dest, ctx)
        return
    raise EngineError(f"Unsupported archive: {path.name}")


def _extract_gzip(path: Path, dest: Path, ctx: Ctx) -> None:
    try:
        with tarfile.open(path, "r:gz") as archive:
            _check_extract_size(sum(member.size for member in archive.getmembers()))
            if ctx.cancelled():
                raise EngineError("Cancelled.")
            _extract_tar(archive, dest)
        return
    except tarfile.TarError:
        pass
    inner_name = path.stem if path.stem else "payload"
    payload_path = dest / inner_name
    total = 0
    with gzip.open(path, "rb") as src, open(payload_path, "wb") as dst:
        head = src.read(_GZ_SNIFF_BYTES)
        dst.write(head)
        total = len(head)
        if total > _MAX_EXTRACT_BYTES:
            raise EngineError("This archive is too large to extract safely.")
        while True:
            if ctx.cancelled():
                raise EngineError("Cancelled.")
            chunk = src.read(1024 * 1024)
            if not chunk:
                break
            total += len(chunk)
            if total > _MAX_EXTRACT_BYTES:
                raise EngineError("This archive is too large to extract safely.")
            dst.write(chunk)
    if _looks_like_tar(head) or (total <= _GZ_SNIFF_BYTES and tarfile.is_tarfile(payload_path)):
        with tarfile.open(payload_path, "r:") as archive:
            if ctx.cancelled():
                raise EngineError("Cancelled.")
            _extract_tar(archive, dest)
        payload_path.unlink(missing_ok=True)


def _create_archive(dest_archive: Path, source_dir: Path, ctx: Ctx) -> None:
    suffix = dest_archive.suffix.lower()
    entries = sorted(source_dir.iterdir())
    if suffix == ".zip":
        with zipfile.ZipFile(dest_archive, "w", zipfile.ZIP_DEFLATED) as archive:
            for entry in entries:
                for file in [entry] if entry.is_file() else entry.rglob("*"):
                    if file.is_file():
                        archive.write(file, file.relative_to(source_dir))
        return
    if suffix == ".tar":
        with tarfile.open(dest_archive, "w:") as archive:
            for entry in entries:
                archive.add(entry, arcname=entry.name)
        return
    if suffix == ".gz":
        with tarfile.open(dest_archive, "w:gz") as archive:
            for entry in entries:
                archive.add(entry, arcname=entry.name)
        return
    if suffix == ".rar":
        rar = rar_path()
        if not rar:
            raise EngineError("WinRAR is required to create RAR archives.")
        names = [
            f".\\{entry.name}" if entry.name.startswith(("-", "@")) else entry.name
            for entry in entries
        ]
        if not names:
            raise EngineError("The archive could not be repackaged.")
        if ctx.cancelled():
            raise EngineError("Cancelled.")
        code = run_process([rar, "a", "-m0", "-ep1", "-y", str(dest_archive), *names], cwd=str(source_dir), cancel=ctx.cancel)
        if code == -9:
            raise EngineError("Cancelled.")
        if code not in (0, 1):
            raise EngineError("Could not create the RAR archive.")
        return
    raise EngineError(f"Unsupported archive target: {dest_archive.name}")


def convert_archive(path: Path, target: str, ctx: Ctx, reserved: set[Path]) -> Path:
    verify_writable(path)
    ctx.status(f"Repackaging {path.name} as {target.upper()}")
    tmp_dir = Path(tempfile.mkdtemp(prefix="tangerine_arc_"))
    extract_dir = tmp_dir / "content"
    extract_dir.mkdir()
    suffix = ".tar.gz" if target == "gz" else f".{target}"
    out_path = _unique(path.parent, path.stem, suffix, reserved)
    try:
        _extract_archive(path, extract_dir, ctx)
        if ctx.cancelled():
            raise EngineError("Cancelled.")
        _create_archive(out_path, extract_dir, ctx)
    except Exception:
        out_path.unlink(missing_ok=True)
        raise
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)
    if not out_path.exists() or out_path.stat().st_size == 0:
        out_path.unlink(missing_ok=True)
        raise EngineError("The archive could not be repackaged.")
    ctx.progress(1.0)
    return out_path


def extract_archive(path: Path, ctx: Ctx, reserved: set[Path]) -> Path:
    verify_writable(path)
    name = path.name
    if name.lower().endswith(".tar.gz"):
        stem = name[: -len(".tar.gz")]
    else:
        stem = path.stem
    directory = path.parent / stem
    index = 2
    while directory.exists():
        directory = path.parent / f"{stem} {index}"
        index += 1
    tmp_dir = Path(tempfile.mkdtemp(prefix=".tangerine_extract_", dir=str(path.parent)))
    ctx.status(f"Extracting {path.name}")
    try:
        _extract_archive(path, tmp_dir, ctx)
        os.replace(tmp_dir, directory)
    except Exception:
        shutil.rmtree(tmp_dir, ignore_errors=True)
        raise
    ctx.progress(1.0)
    return directory


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------

def convert_single(path: Path, target: str, ctx: Ctx, reserved: set[Path]) -> list[Path]:
    if not path.exists():
        raise EngineError(f"The source file could not be found: {path.name}")
    family = family_of(path)
    if family == FAMILY_IMAGE:
        if target == "docx":
            return [image_to_docx(path, ctx, reserved)]
        return [convert_image(path, target, ctx, reserved)]
    if family == FAMILY_AUDIO:
        return [convert_audio(path, target, ctx, reserved)]
    if family == FAMILY_VIDEO:
        return [convert_video(path, target, ctx, reserved)]
    if family == FAMILY_GIF:
        return [convert_gif(path, target, ctx, reserved)]
    if family == FAMILY_PDF:
        return convert_pdf(path, target, ctx, reserved)
    if family == FAMILY_TXT:
        return convert_txt(path, target, ctx, reserved)
    if family == FAMILY_DOC:
        from . import documents

        return documents.convert_document(path, target, ctx, reserved)
    if family == FAMILY_ARCHIVE:
        return [convert_archive(path, target, ctx, reserved)]
    raise EngineError(f"No conversion engine for {path.name}")
